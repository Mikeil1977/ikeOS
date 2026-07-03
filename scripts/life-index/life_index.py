from __future__ import annotations

import argparse
import csv
import json
import os
import re
import sqlite3
import sys
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable


DEFAULT_DB_PATH = Path(r"C:\IkeOS\private\life-index\life-index.sqlite")

DEFAULT_EXTENSIONS = {
    ".pdf",
    ".doc",
    ".docx",
    ".xls",
    ".xlsx",
    ".csv",
    ".txt",
    ".md",
    ".rtf",
    ".jpg",
    ".jpeg",
    ".png",
    ".heic",
    ".zip",
    ".7z",
}

TEXT_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".csv", ".txt", ".md", ".rtf"}
MEDIA_EXTENSIONS = {".jpg", ".jpeg", ".png", ".heic"}
ARCHIVE_EXTENSIONS = {".zip", ".7z"}

EXCLUDED_COMPONENTS = {
    "$recycle.bin",
    "$windows.~bt",
    "$windows.~ws",
    "$winreagent",
    ".cache",
    ".codex",
    ".config",
    "codex",
    ".docker",
    ".git",
    ".hg",
    ".mypy_cache",
    ".nuget",
    ".ssh",
    ".svn",
    ".tox",
    ".venv",
    ".vscode",
    "__pycache__",
    "appdata",
    "application data",
    "build",
    "dist",
    "local settings",
    "node_modules",
    "onedrivecloudtemp",
    "packages",
    "program files",
    "program files (x86)",
    "programdata",
    "recovery",
    "steamapps",
    "system volume information",
    "temp",
    "temporary internet files",
    "tmp",
    "venv",
    "windows",
    "windowsapps",
    "xboxgames",
}


@dataclass(frozen=True)
class RootSpec:
    label: str
    path: Path


DEFAULT_ROOT_SPECS = [
    RootSpec("Desktop", Path(r"C:\Users\MikeDalziel\Desktop")),
    RootSpec("Documents", Path(r"C:\Users\MikeDalziel\Documents")),
    RootSpec("Downloads", Path(r"C:\Users\MikeDalziel\Downloads")),
    RootSpec("OneDrive personal", Path(r"C:\Users\MikeDalziel\OneDrive")),
    RootSpec(
        "OneDrive Invigorate",
        Path(r"C:\Users\MikeDalziel\OneDrive - Invigorate IT Ltd"),
    ),
    RootSpec("Elysium Dynamics sync", Path(r"C:\Users\MikeDalziel\Elysium Dynamics Ltd")),
    RootSpec("Invigorate IT sync", Path(r"C:\Users\MikeDalziel\Invigorate IT Ltd")),
    RootSpec("D Users archive", Path(r"D:\Users")),
    RootSpec("D ASB", Path(r"D:\ASB")),
    RootSpec("D InterAccount", Path(r"D:\InterAccount")),
    RootSpec("D Server Files", Path(r"D:\Server Files")),
    RootSpec("Local Reference", Path(r"C:\Local Reference")),
]


CATEGORY_KEYWORDS = [
    ("insurance", ("insurance", "policy", "renewal", "claim", "contents", "building")),
    ("finance/tax", ("invoice", "receipt", "tax", "vat", "hmrc", "bank", "pension", "payroll")),
    ("home/property", ("mortgage", "property", "tenancy", "council", "utility", "water", "gas", "electric")),
    ("vehicle", ("vehicle", "car", "mot", "dvla", "v5c", "garage")),
    ("identity", ("passport", "licence", "birth certificate", "identity", "id check")),
    ("legal", ("legal", "contract", "agreement", "solicitor", "will", "terms")),
    ("health", ("health", "medical", "gp", "nhs", "clinic", "prescription", "sleep study")),
    ("education/family", ("school", "education", "homework", "family", "ella", "sarah")),
    ("work/customer", ("elysium", "invigorate", "enermech", "twma", "customer", "project")),
]


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def normalize_path(path: Path) -> str:
    return str(path.resolve(strict=False))


def existing_default_roots() -> list[Path]:
    return [spec.path for spec in DEFAULT_ROOT_SPECS if spec.path.exists()]


def should_exclude(path: Path, root: Path | None = None) -> bool:
    if root is not None:
        try:
            parts_source = path.relative_to(root).parts
        except ValueError:
            parts_source = path.parts
    else:
        parts_source = path.parts
    parts = [part.lower() for part in parts_source if part not in {"", "."}]
    for part in parts:
        if part in EXCLUDED_COMPONENTS:
            return True
        if part.startswith(".") and part not in {".", ".."}:
            return True
    return False


def iter_candidate_files(roots: Iterable[Path], extensions: set[str] | None = None) -> Iterable[tuple[Path, Path]]:
    allowed = extensions or DEFAULT_EXTENSIONS
    for root in roots:
        if not root.exists():
            continue
        stack = [root]
        while stack:
            current = stack.pop()
            if should_exclude(current, root=root) and current != root:
                continue
            try:
                children = list(current.iterdir())
            except OSError:
                continue
            for child in children:
                if should_exclude(child, root=root):
                    continue
                if child.is_dir():
                    stack.append(child)
                    continue
                if child.is_file() and child.suffix.lower() in allowed:
                    yield root, child


def open_database(db_path: Path) -> sqlite3.Connection:
    db_path.parent.mkdir(parents=True, exist_ok=True)
    con = sqlite3.connect(db_path)
    con.row_factory = sqlite3.Row
    con.execute("PRAGMA journal_mode=WAL")
    con.execute("PRAGMA synchronous=NORMAL")
    initialize_schema(con)
    return con


def initialize_schema(con: sqlite3.Connection) -> None:
    con.executescript(
        """
        CREATE TABLE IF NOT EXISTS scan_runs (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            started_at TEXT NOT NULL,
            finished_at TEXT,
            mode TEXT NOT NULL,
            roots_json TEXT NOT NULL,
            extract_text INTEGER NOT NULL,
            dry_run INTEGER NOT NULL,
            status TEXT NOT NULL,
            files_seen INTEGER NOT NULL DEFAULT 0,
            files_indexed INTEGER NOT NULL DEFAULT 0,
            errors INTEGER NOT NULL DEFAULT 0
        );

        CREATE TABLE IF NOT EXISTS roots (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            path TEXT NOT NULL UNIQUE,
            label TEXT,
            last_seen_at TEXT NOT NULL
        );

        CREATE TABLE IF NOT EXISTS files (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            path TEXT NOT NULL UNIQUE,
            root_path TEXT NOT NULL,
            relative_path TEXT NOT NULL,
            name TEXT NOT NULL,
            extension TEXT NOT NULL,
            size_bytes INTEGER NOT NULL,
            mtime_utc TEXT NOT NULL,
            ctime_utc TEXT NOT NULL,
            category TEXT NOT NULL,
            scan_status TEXT NOT NULL,
            text_extracted INTEGER NOT NULL,
            extraction_error TEXT,
            is_media INTEGER NOT NULL,
            updated_at TEXT NOT NULL
        );

        CREATE TABLE IF NOT EXISTS extraction_errors (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            file_id INTEGER,
            path TEXT NOT NULL,
            error TEXT NOT NULL,
            created_at TEXT NOT NULL
        );

        CREATE TABLE IF NOT EXISTS image_metadata (
            file_id INTEGER PRIMARY KEY,
            path TEXT NOT NULL UNIQUE,
            width INTEGER,
            height INTEGER,
            image_format TEXT,
            updated_at TEXT NOT NULL
        );

        CREATE VIRTUAL TABLE IF NOT EXISTS file_text_fts
        USING fts5(path UNINDEXED, name, category UNINDEXED, text);
        """
    )
    con.commit()


def guess_category(path: Path, text: str = "") -> str:
    blob = f"{path.name} {path.parent} {text[:4000]}".lower()
    if path.suffix.lower() in MEDIA_EXTENSIONS:
        return "media/photo"
    if path.suffix.lower() in ARCHIVE_EXTENSIONS:
        return "archive"
    for category, keywords in CATEGORY_KEYWORDS:
        if any(keyword in blob for keyword in keywords):
            return category
    return "unknown"


def stat_times(path: Path) -> tuple[str, str]:
    stat = path.stat()
    mtime = datetime.fromtimestamp(stat.st_mtime, timezone.utc).isoformat(timespec="seconds")
    ctime = datetime.fromtimestamp(stat.st_ctime, timezone.utc).isoformat(timespec="seconds")
    return mtime, ctime


def extract_text(path: Path, max_chars: int = 200_000) -> str:
    suffix = path.suffix.lower()
    if suffix not in TEXT_EXTENSIONS:
        return ""
    if suffix in {".txt", ".md", ".csv", ".rtf"}:
        return extract_plain_text(path, max_chars=max_chars)
    if suffix == ".pdf":
        return extract_pdf_text(path, max_chars=max_chars)
    if suffix == ".docx":
        return extract_docx_text(path, max_chars=max_chars)
    if suffix == ".xlsx":
        return extract_xlsx_text(path, max_chars=max_chars)
    return ""


def extract_plain_text(path: Path, max_chars: int) -> str:
    raw = path.read_bytes()[: max_chars * 4]
    for encoding in ("utf-8", "utf-16", "cp1252"):
        try:
            return raw.decode(encoding, errors="ignore")[:max_chars]
        except UnicodeError:
            continue
    return raw.decode("utf-8", errors="ignore")[:max_chars]


def extract_pdf_text(path: Path, max_chars: int) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    parts: list[str] = []
    for page in reader.pages[:50]:
        parts.append(page.extract_text() or "")
        if sum(len(part) for part in parts) >= max_chars:
            break
    return "\n".join(parts)[:max_chars]


def extract_docx_text(path: Path, max_chars: int) -> str:
    import docx

    document = docx.Document(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs]
    return "\n".join(parts)[:max_chars]


def extract_xlsx_text(path: Path, max_chars: int) -> str:
    import openpyxl

    workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    parts: list[str] = []
    try:
        for sheet in workbook.worksheets[:10]:
            parts.append(sheet.title)
            for row in sheet.iter_rows(max_row=300, values_only=True):
                values = [str(value) for value in row if value not in (None, "")]
                if values:
                    parts.append(" | ".join(values))
                if sum(len(part) for part in parts) >= max_chars:
                    return "\n".join(parts)[:max_chars]
    finally:
        workbook.close()
    return "\n".join(parts)[:max_chars]


def read_image_metadata(path: Path) -> dict[str, object]:
    from PIL import Image

    with Image.open(path) as image:
        width, height = image.size
        return {
            "width": width,
            "height": height,
            "image_format": image.format or "",
        }


def upsert_root(con: sqlite3.Connection, root: Path) -> None:
    root_path = normalize_path(root)
    label = next((spec.label for spec in DEFAULT_ROOT_SPECS if normalize_path(spec.path) == root_path), root.name)
    con.execute(
        """
        INSERT INTO roots(path, label, last_seen_at)
        VALUES (?, ?, ?)
        ON CONFLICT(path) DO UPDATE SET label=excluded.label, last_seen_at=excluded.last_seen_at
        """,
        (root_path, label, utc_now()),
    )


def index_file(con: sqlite3.Connection, root: Path, file_path: Path, extract: bool) -> tuple[bool, str | None]:
    path = normalize_path(file_path)
    existing = con.execute(
        "SELECT category, text_extracted, extraction_error FROM files WHERE path = ?",
        (path,),
    ).fetchone()
    text = ""
    extraction_error: str | None = None
    metadata_error: str | None = None
    image_metadata: dict[str, object] | None = None
    text_extracted = False
    is_media = file_path.suffix.lower() in MEDIA_EXTENSIONS
    try:
        if extract:
            text = extract_text(file_path)
            text_extracted = bool(text)
    except Exception as exc:  # noqa: BLE001 - indexing must not fail on one bad file.
        extraction_error = f"{type(exc).__name__}: {exc}"

    try:
        mtime_utc, ctime_utc = stat_times(file_path)
        size = file_path.stat().st_size
    except OSError as exc:
        return False, f"{type(exc).__name__}: {exc}"

    if not extract and existing is not None:
        category = existing["category"]
        text_extracted = bool(existing["text_extracted"])
        extraction_error = existing["extraction_error"]
    else:
        category = guess_category(file_path, text)
    if is_media:
        try:
            image_metadata = read_image_metadata(file_path)
        except Exception as exc:  # noqa: BLE001 - keep indexing even if one image is odd.
            metadata_error = f"{type(exc).__name__}: {exc}"
    root_path = normalize_path(root)
    try:
        relative_path = str(file_path.relative_to(root))
    except ValueError:
        relative_path = file_path.name

    con.execute(
        """
        INSERT INTO files(
            path, root_path, relative_path, name, extension, size_bytes,
            mtime_utc, ctime_utc, category, scan_status, text_extracted,
            extraction_error, is_media, updated_at
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ON CONFLICT(path) DO UPDATE SET
            root_path=excluded.root_path,
            relative_path=excluded.relative_path,
            name=excluded.name,
            extension=excluded.extension,
            size_bytes=excluded.size_bytes,
            mtime_utc=excluded.mtime_utc,
            ctime_utc=excluded.ctime_utc,
            category=excluded.category,
            scan_status=excluded.scan_status,
            text_extracted=excluded.text_extracted,
            extraction_error=excluded.extraction_error,
            is_media=excluded.is_media,
            updated_at=excluded.updated_at
        """,
        (
            path,
            root_path,
            relative_path,
            file_path.name,
            file_path.suffix.lower(),
            size,
            mtime_utc,
            ctime_utc,
            category,
            "indexed",
            1 if text_extracted else 0,
            extraction_error,
            1 if is_media else 0,
            utc_now(),
        ),
    )

    if extract and text:
        con.execute("DELETE FROM file_text_fts WHERE path = ?", (path,))
        con.execute(
            "INSERT INTO file_text_fts(path, name, category, text) VALUES (?, ?, ?, ?)",
            (path, file_path.name, category, text),
        )

    file_id = con.execute("SELECT id FROM files WHERE path = ?", (path,)).fetchone()["id"]

    if image_metadata:
        con.execute(
            """
            INSERT INTO image_metadata(file_id, path, width, height, image_format, updated_at)
            VALUES (?, ?, ?, ?, ?, ?)
            ON CONFLICT(file_id) DO UPDATE SET
                path=excluded.path,
                width=excluded.width,
                height=excluded.height,
                image_format=excluded.image_format,
                updated_at=excluded.updated_at
            """,
            (
                file_id,
                path,
                image_metadata["width"],
                image_metadata["height"],
                image_metadata["image_format"],
                utc_now(),
            ),
        )

    if extraction_error or metadata_error:
        con.execute(
            "INSERT INTO extraction_errors(file_id, path, error, created_at) VALUES (?, ?, ?, ?)",
            (file_id, path, extraction_error or metadata_error or "", utc_now()),
        )

    return True, extraction_error or metadata_error


def build_index(
    db_path: Path = DEFAULT_DB_PATH,
    roots: Iterable[Path] | None = None,
    dry_run: bool = False,
    extract_text: bool = True,
    no_text: bool = False,
) -> dict[str, object]:
    selected_roots = [Path(root) for root in (roots if roots is not None else existing_default_roots())]
    selected_roots = [root for root in selected_roots if root.exists()]
    result: dict[str, object] = {
        "mode": "dry_run" if dry_run else "scan",
        "db_path": str(db_path),
        "existing_roots": [str(root) for root in selected_roots],
        "candidate_extensions": sorted(DEFAULT_EXTENSIONS),
        "excluded_path_fragments": sorted(EXCLUDED_COMPONENTS),
        "files_seen": 0,
        "files_indexed": 0,
        "errors": 0,
    }
    if dry_run:
        return result

    con = open_database(db_path)
    run_id = con.execute(
        """
        INSERT INTO scan_runs(started_at, mode, roots_json, extract_text, dry_run, status)
        VALUES (?, ?, ?, ?, ?, ?)
        """,
        (
            utc_now(),
            "scan",
            json.dumps([str(root) for root in selected_roots], ensure_ascii=False),
            0 if no_text else int(extract_text),
            0,
            "running",
        ),
    ).lastrowid

    files_seen = 0
    files_indexed = 0
    errors = 0
    status = "failed"
    try:
        for root in selected_roots:
            upsert_root(con, root)
            root_path_key = normalize_path(root)
            con.execute(
                "UPDATE files SET scan_status = 'pending_rescan', updated_at = ? WHERE root_path = ?",
                (utc_now(), root_path_key),
            )
            for root_path, file_path in iter_candidate_files([root]):
                files_seen += 1
                ok, error = index_file(con, root_path, file_path, extract=extract_text and not no_text)
                if ok:
                    files_indexed += 1
                if error:
                    errors += 1
                if files_seen % 250 == 0:
                    con.commit()
            con.execute(
                "UPDATE files SET scan_status = 'stale', updated_at = ? WHERE root_path = ? AND scan_status = 'pending_rescan'",
                (utc_now(), root_path_key),
            )
        status = "complete"
    except Exception:  # noqa: BLE001
        status = "failed"
        raise
    finally:
        con.execute(
            """
            UPDATE scan_runs
            SET finished_at=?, status=?, files_seen=?, files_indexed=?, errors=?
            WHERE id=?
            """,
            (utc_now(), status, files_seen, files_indexed, errors, run_id),
        )
        con.commit()
        con.close()

    result["files_seen"] = files_seen
    result["files_indexed"] = files_indexed
    result["errors"] = errors
    return result


def query_index(
    db_path: Path = DEFAULT_DB_PATH,
    search: str | None = None,
    category: str | None = None,
    root: str | None = None,
    since: str | None = None,
    until: str | None = None,
    limit: int = 25,
) -> list[dict[str, object]]:
    con = sqlite3.connect(db_path)
    con.row_factory = sqlite3.Row
    try:
        clauses: list[str] = []
        params: list[object] = []
        if category:
            clauses.append("f.category = ?")
            params.append(category)
        if root:
            clauses.append("f.root_path LIKE ?")
            params.append(f"%{root}%")
        if since:
            clauses.append("f.mtime_utc >= ?")
            params.append(normalize_date_bound(since))
        if until:
            clauses.append("f.mtime_utc <= ?")
            params.append(normalize_date_bound(until, end_of_day=True))
        clauses.append("f.scan_status = 'indexed'")
        if search:
            clauses.append(
                """
                (
                    f.name LIKE ?
                    OR f.relative_path LIKE ?
                    OR f.path IN (
                        SELECT path FROM file_text_fts WHERE file_text_fts MATCH ?
                    )
                )
                """
            )
            params.extend([f"%{search}%", f"%{search}%", fts_literal_query(search)])
        where = "WHERE " + " AND ".join(clauses)
        rows = con.execute(
            f"""
            SELECT f.path, f.root_path, f.relative_path, f.name, f.extension,
                   f.size_bytes, f.mtime_utc, f.category, f.scan_status,
                   f.text_extracted, f.extraction_error
            FROM files f
            {where}
            ORDER BY f.mtime_utc DESC
            LIMIT ?
            """,
            [*params, limit],
        ).fetchall()
        return [dict(row) for row in rows]
    finally:
        con.close()


def category_summary(con: sqlite3.Connection) -> list[sqlite3.Row]:
    return con.execute(
        """
        SELECT category, COUNT(*) AS count, MAX(mtime_utc) AS newest
        FROM files
        WHERE scan_status = 'indexed'
        GROUP BY category
        ORDER BY count DESC, category ASC
        """
    ).fetchall()


def top_roots_for_category(con: sqlite3.Connection, category: str, limit: int = 3) -> list[str]:
    rows = con.execute(
        """
        SELECT root_path, COUNT(*) AS count
        FROM files
        WHERE category = ? AND scan_status = 'indexed'
        GROUP BY root_path
        ORDER BY count DESC
        LIMIT ?
        """,
        (category, limit),
    ).fetchall()
    return [f"{Path(row['root_path']).name or row['root_path']} ({row['count']})" for row in rows]


def representative_names(con: sqlite3.Connection, category: str, limit: int = 5) -> list[str]:
    rows = con.execute(
        """
        SELECT name
        FROM files
        WHERE category = ? AND scan_status = 'indexed'
        ORDER BY mtime_utc DESC
        LIMIT ?
        """,
        (category, limit),
    ).fetchall()
    return [row["name"] for row in rows]


def render_safe_report(db_path: Path = DEFAULT_DB_PATH, include_filenames: bool = False) -> str:
    con = sqlite3.connect(db_path)
    con.row_factory = sqlite3.Row
    try:
        total = con.execute("SELECT COUNT(*) AS c FROM files WHERE scan_status = 'indexed'").fetchone()["c"]
        lines = [
            "# MikeOS Life Index Summary",
            "",
            f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M %z')}".rstrip(),
            "",
            f"Indexed files: {total}",
            "",
            "| Category | Count | Newest item | Top roots |",
            "| --- | ---: | --- | --- |",
        ]
        for row in category_summary(con):
            roots = ", ".join(top_roots_for_category(con, row["category"]))
            lines.append(f"| {row['category']} | {row['count']} | {row['newest'] or ''} | {roots} |")
            if include_filenames:
                names = ", ".join(representative_names(con, row["category"]))
                if names:
                    lines.append(f"| {row['category']} representative filenames |  |  | {names} |")
        lines.extend(
            [
                "",
                "Privacy note: this report is distilled metadata only. It does not include extracted document text.",
            ]
        )
        return "\n".join(lines)
    finally:
        con.close()


def normalize_date_bound(value: str, end_of_day: bool = False) -> str:
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
        return f"{value}T{'23:59:59' if end_of_day else '00:00:00'}+00:00"
    return value


def fts_literal_query(value: str) -> str:
    escaped = value.replace('"', '""')
    return f'"{escaped}"'


def write_report(db_path: Path, output: Path, include_filenames: bool = False) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(render_safe_report(db_path=db_path, include_filenames=include_filenames), encoding="utf-8")


def parse_roots(raw_roots: list[str] | None) -> list[Path] | None:
    if not raw_roots:
        return None
    roots: list[Path] = []
    for raw in raw_roots:
        for piece in raw.split(";"):
            piece = piece.strip()
            if piece:
                roots.append(Path(piece))
    return roots


def command_build(args: argparse.Namespace) -> int:
    roots = parse_roots(args.root)
    result = build_index(
        db_path=Path(args.db),
        roots=roots,
        dry_run=args.dry_run,
        extract_text=not args.no_text,
        no_text=args.no_text,
    )
    print(json.dumps(result, indent=2, ensure_ascii=False))
    return 0


def command_query(args: argparse.Namespace) -> int:
    matches = query_index(
        db_path=Path(args.db),
        search=args.search,
        category=args.category,
        root=args.root,
        since=args.since,
        until=args.until,
        limit=args.limit,
    )
    if args.format == "json":
        print(json.dumps(matches, indent=2, ensure_ascii=False))
        return 0
    writer = csv.DictWriter(sys.stdout, fieldnames=list(matches[0].keys()) if matches else ["path"])
    writer.writeheader()
    writer.writerows(matches)
    return 0


def command_report(args: argparse.Namespace) -> int:
    report = render_safe_report(db_path=Path(args.db), include_filenames=args.include_filenames)
    if args.output:
        output = Path(args.output)
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_text(report, encoding="utf-8")
    print(report)
    return 0


def make_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="MikeOS private life index")
    parser.add_argument("--db", default=str(DEFAULT_DB_PATH), help="Private SQLite DB path")
    subparsers = parser.add_subparsers(dest="command", required=True)

    build = subparsers.add_parser("build", help="Build or update the private life index")
    build.add_argument("--root", action="append", help="Root path to scan. Repeat or separate with semicolons.")
    build.add_argument("--dry-run", action="store_true", help="List roots and exclusions without writing the DB")
    build.add_argument("--resume", action="store_true", help="Resume by safely rerunning selected roots")
    build.add_argument("--no-text", action="store_true", help="Index metadata only; skip text extraction")
    build.set_defaults(func=command_build)

    query = subparsers.add_parser("query", help="Search the private life index")
    query.add_argument("--search", help="Search term for filename/path/text")
    query.add_argument("--category", help="Category filter")
    query.add_argument("--root", help="Root path/name filter")
    query.add_argument("--since", help="Only include files modified on or after this date or timestamp")
    query.add_argument("--until", help="Only include files modified on or before this date or timestamp")
    query.add_argument("--limit", type=int, default=25)
    query.add_argument("--format", choices=["csv", "json"], default="csv")
    query.set_defaults(func=command_query)

    report = subparsers.add_parser("report", help="Render a MikeOS-safe summary report")
    report.add_argument("--output", help="Optional report output path")
    report.add_argument("--include-filenames", action="store_true", help="Include representative filenames only")
    report.set_defaults(func=command_report)

    return parser


def main(argv: list[str] | None = None) -> int:
    parser = make_parser()
    args = parser.parse_args(argv)
    return args.func(args)


if __name__ == "__main__":
    raise SystemExit(main())
